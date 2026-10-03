#!/usr/bin/env python3
from pathlib import Path

import pandas as pd


OUTDIR = Path("path/to/project/10.centromere_analysis/output_synthen/06_functional_centromere_sequence_composition/R108_TEsorter_TRASH_test")
LONG = OUTDIR / "genome_R108.functional_centromere_sequence_composition.broad.long.tsv"

CATEGORY_ORDER = [
    "Tandem repeat",
    "LTR/Gypsy",
    "LTR/Copia",
    "LTR unknown",
    "SINE / LINE",
    "DNA transposons",
    "Unclassified TE",
    "Others",
]

CATEGORY_DISPLAY = {
    "Tandem repeat": "Tandem repeat",
    "LTR/Gypsy": "LTR/Gypsy",
    "LTR/Copia": "LTR/Copia",
    "LTR unknown": "LTR unknown",
    "SINE / LINE": "SINE / LINE",
    "DNA transposons": "DNA transposons",
    "Unclassified TE": "Unclassified TE",
    "Others": "Others",
}


def make_centromere_name(value):
    # R108.CEN01 -> Mtr_R108.CEN01
    suffix = str(value).split(".", 1)[1] if "." in str(value) else str(value)
    return f"Mtr_R108.{suffix}"


def build_tables():
    df = pd.read_csv(LONG, sep="\t")
    df["Centromere"] = df["centromere"].map(make_centromere_name)
    df["chr_num"] = df["chr"].str.replace("Chr", "", regex=False).astype(int)

    rows = []
    for _, sub in df.groupby(["chr_num", "Centromere"], sort=True):
        sub = sub.set_index("category")
        row = {"Centromere": sub["Centromere"].iloc[0]}
        for category in CATEGORY_ORDER:
            label = CATEGORY_DISPLAY[category]
            if category in sub.index:
                row[f"{label} length (bp)"] = int(round(float(sub.loc[category, "length_bp"])))
                row[f"{label} ratio (%)"] = round(float(sub.loc[category, "ratio_pct"]), 2)
            else:
                row[f"{label} length (bp)"] = 0
                row[f"{label} ratio (%)"] = 0.0
        rows.append(row)

    table = pd.DataFrame(rows)
    avg = {"Centromere": "Average"}
    for category in CATEGORY_ORDER:
        label = CATEGORY_DISPLAY[category]
        avg[f"{label} length (bp)"] = int(round(table[f"{label} length (bp)"].mean()))
        avg[f"{label} ratio (%)"] = round(float(table[f"{label} ratio (%)"].mean()), 2)
    table = pd.concat([table, pd.DataFrame([avg])], ignore_index=True)
    return table


def write_two_row_header_tsv(table, path):
    first = ["Centromere"]
    second = [""]
    for category in CATEGORY_ORDER:
        label = CATEGORY_DISPLAY[category]
        first.extend([label, ""])
        second.extend(["length (bp)", "ratio (%)"])
    with open(path, "w", newline="") as handle:
        handle.write("\t".join(first) + "\n")
        handle.write("\t".join(second) + "\n")
        for _, row in table.iterrows():
            values = [row["Centromere"]]
            for category in CATEGORY_ORDER:
                label = CATEGORY_DISPLAY[category]
                values.append(str(int(row[f"{label} length (bp)"])))
                values.append(f"{float(row[f'{label} ratio (%)']):.2f}")
            handle.write("\t".join(values) + "\n")


def write_xlsx_if_possible(table, path):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        return f"Skipped xlsx: {exc}"

    wb = Workbook()
    ws = wb.active
    ws.title = "R108_TableS11"
    ws.cell(1, 1, "Table S11. Sequence composition within functional centromeric regions in genome_R108.")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=1 + len(CATEGORY_ORDER) * 2)
    ws.cell(2, 1, "Centromere")
    ws.merge_cells(start_row=2, start_column=1, end_row=3, end_column=1)

    col = 2
    for category in CATEGORY_ORDER:
        label = CATEGORY_DISPLAY[category]
        ws.cell(2, col, label)
        ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)
        ws.cell(3, col, "length (bp)")
        ws.cell(3, col + 1, "ratio (%)")
        col += 2

    start_row = 4
    for r, (_, row) in enumerate(table.iterrows(), start=start_row):
        ws.cell(r, 1, row["Centromere"])
        col = 2
        for category in CATEGORY_ORDER:
            label = CATEGORY_DISPLAY[category]
            ws.cell(r, col, int(row[f"{label} length (bp)"]))
            ws.cell(r, col + 1, float(row[f"{label} ratio (%)"]))
            col += 2

    thin = Side(style="thin", color="D9D9D9")
    header_fill = PatternFill("solid", fgColor="F2F2F2")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
            if cell.row <= 3:
                cell.font = Font(bold=True)
                cell.fill = header_fill
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True)

    ws.freeze_panes = "B4"
    widths = {1: 18}
    for idx in range(2, ws.max_column + 1):
        widths[idx] = 14 if (idx - 2) % 2 == 0 else 10
    for idx, width in widths.items():
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in range(start_row, ws.max_row + 1):
        for col in range(2, ws.max_column + 1, 2):
            ws.cell(row, col).number_format = "#,##0"
            ws.cell(row, col + 1).number_format = "0.00"
    wb.save(path)
    return "Wrote xlsx"


def main():
    table = build_tables()
    flat = OUTDIR / "genome_R108.functional_centromere_sequence_composition.TableS11_style.flat.tsv"
    pretty = OUTDIR / "genome_R108.functional_centromere_sequence_composition.TableS11_style.tsv"
    xlsx = OUTDIR / "genome_R108.functional_centromere_sequence_composition.TableS11_style.xlsx"
    table.to_csv(flat, sep="\t", index=False)
    write_two_row_header_tsv(table, pretty)
    status = write_xlsx_if_possible(table, xlsx)
    print(f"Wrote: {flat}")
    print(f"Wrote: {pretty}")
    print(status)
    print(table.to_string(index=False))


if __name__ == "__main__":
    main()
